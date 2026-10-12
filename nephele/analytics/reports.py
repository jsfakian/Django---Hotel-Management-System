"""
Custom report data sets and file rendering.

``build_report`` returns ``{"columns", "rows", "summary"}`` for a CustomReport;
``render_report`` turns that into (bytes, content_type, extension).
"""

import csv
import io
from collections import defaultdict
from decimal import Decimal

from django.db.models import Avg, Count
from django.utils import timezone

from properties.models import Property
from room.models import Booking

from .metrics import REVENUE_STATUSES, SOURCE_LABELS, NightLedger, money, pct


def _cols(*pairs):
    return [{"field": field, "header": header} for field, header in pairs]


def _guest_name(guest):
    return f"{guest.first_name} {guest.last_name}".strip() if guest else ""


def _revenue(properties, report):
    ledger = NightLedger(properties, report.from_date, report.to_date)
    rows = [
        {
            "date": r["date"],
            "revenue": r["revenue"],
            "rooms_sold": r["rooms_sold"],
            "occupancy": r["occupancy"],
            "adr": r["adr"],
            "revpar": r["revpar"],
            "arrivals": r["bookings"],
        }
        for r in ledger.series()
    ]
    k = ledger.kpis()
    return (
        _cols(
            ("date", "Date"),
            ("revenue", "Room revenue (€)"),
            ("rooms_sold", "Rooms sold"),
            ("occupancy", "Occupancy %"),
            ("adr", "ADR (€)"),
            ("revpar", "RevPAR (€)"),
            ("arrivals", "Arrivals"),
        ),
        rows,
        {
            "Room revenue (€)": k["total_revenue"],
            "Occupancy %": k["occupancy_rate"],
            "ADR (€)": k["adr"],
            "RevPAR (€)": k["revpar"],
        },
    )


def _occupancy(properties, report):
    ledger = NightLedger(properties, report.from_date, report.to_date)
    departures = defaultdict(int)
    for b in ledger.bookings:
        departures[b.check_out_date.isoformat()] += 1
    rows = [
        {
            "date": r["date"],
            "rooms_available": ledger.rooms,
            "rooms_sold": r["rooms_sold"],
            "occupancy": r["occupancy"],
            "arrivals": r["bookings"],
            "departures": departures[r["date"]],
        }
        for r in ledger.series()
    ]
    k = ledger.kpis()
    return (
        _cols(
            ("date", "Date"),
            ("rooms_available", "Rooms available"),
            ("rooms_sold", "Rooms sold"),
            ("occupancy", "Occupancy %"),
            ("arrivals", "Arrivals"),
            ("departures", "Departures"),
        ),
        rows,
        {
            "Average occupancy %": k["occupancy_rate"],
            "Room-nights sold": k["rooms_sold"],
        },
    )


def _bookings(properties, report):
    qs = (
        Booking.objects.filter(
            room__property__in=properties,
            check_in_date__gte=report.from_date,
            check_in_date__lte=report.to_date,
        )
        .select_related("guest", "room__property", "rate_plan")
        .order_by("check_in_date", "id")
    )
    rows = [
        {
            "booking": b.id,
            "guest": _guest_name(b.guest),
            "property": b.room.property.name if b.room else "",
            "room": b.room.room_number if b.room else "",
            "check_in": b.check_in_date.isoformat(),
            "check_out": b.check_out_date.isoformat(),
            "nights": (b.check_out_date - b.check_in_date).days,
            "guests": b.number_of_guests,
            "status": b.status,
            "source": SOURCE_LABELS.get(b.booking_source, b.booking_source),
            "rate_plan": b.rate_plan.name if b.rate_plan else "",
            "total": money(b.actual_price or b.base_price),
        }
        for b in qs
    ]
    active = [r for r in rows if r["status"] in REVENUE_STATUSES]
    return (
        _cols(
            ("booking", "Booking #"),
            ("guest", "Guest"),
            ("property", "Property"),
            ("room", "Room"),
            ("check_in", "Check-in"),
            ("check_out", "Check-out"),
            ("nights", "Nights"),
            ("guests", "Guests"),
            ("status", "Status"),
            ("source", "Source"),
            ("rate_plan", "Rate plan"),
            ("total", "Total (€)"),
        ),
        rows,
        {
            "Bookings": len(rows),
            "Cancelled": sum(r["status"] == "cancelled" for r in rows),
            "Booked value (€)": money(sum(r["total"] for r in active)),
        },
    )


def _payments(properties, report):
    from payments.models import Payment

    qs = (
        Payment.objects.filter(
            booking__room__property__in=properties,
            created_at__date__gte=report.from_date,
            created_at__date__lte=report.to_date,
        )
        .select_related("guest", "payment_method", "booking")
        .order_by("created_at")
    )
    rows = [
        {
            "reference": p.reference_code,
            "date": timezone.localtime(p.created_at).strftime("%Y-%m-%d"),
            "guest": _guest_name(p.guest),
            "booking": p.booking_id or "",
            "method": p.payment_method.name,
            "status": p.status,
            "amount": money(p.amount),
            "currency": p.currency,
        }
        for p in qs
    ]

    def total(status):
        return money(sum(r["amount"] for r in rows if r["status"] in status))

    return (
        _cols(
            ("reference", "Reference"),
            ("date", "Date"),
            ("guest", "Guest"),
            ("booking", "Booking #"),
            ("method", "Method"),
            ("status", "Status"),
            ("amount", "Amount"),
            ("currency", "Currency"),
        ),
        rows,
        {
            "Payments": len(rows),
            "Collected (€)": total({"completed"}),
            "Outstanding (€)": total({"pending", "processing"}),
            "Refunded (€)": total({"refunded"}),
        },
    )


def _guests(properties, report):
    ledger = NightLedger(properties, report.from_date, report.to_date)
    per_guest = defaultdict(
        lambda: {"stays": set(), "nights": 0, "revenue": Decimal("0")}
    )
    for _, booking, rate in ledger.nights:
        if booking.guest_id:
            entry = per_guest[booking.guest]
            entry["stays"].add(booking.id)
            entry["nights"] += 1
            entry["revenue"] += rate
    rows = sorted(
        [
            {
                "guest": _guest_name(g),
                "email": g.email,
                "country": g.country or g.nationality or "",
                "stays": len(v["stays"]),
                "nights": v["nights"],
                "revenue": money(v["revenue"]),
                "lifetime_bookings": g.number_of_bookings,
            }
            for g, v in per_guest.items()
        ],
        key=lambda r: r["revenue"],
        reverse=True,
    )
    return (
        _cols(
            ("guest", "Guest"),
            ("email", "Email"),
            ("country", "Country"),
            ("stays", "Stays"),
            ("nights", "Nights"),
            ("revenue", "Room revenue (€)"),
            ("lifetime_bookings", "Lifetime bookings"),
        ),
        rows,
        {
            "Guests": len(rows),
            "Repeat guests": sum(r["lifetime_bookings"] > 1 for r in rows),
        },
    )


def _pricing(properties, report):
    from bookings.models import PricingHistory

    qs = (
        PricingHistory.objects.filter(
            room__property__in=properties,
            date__gte=report.from_date,
            date__lte=report.to_date,
        )
        .values("date")
        .annotate(
            base=Avg("base_price"),
            dynamic=Avg("dynamic_price"),
            competitor=Avg("competitor_price"),
            occupancy=Avg("occupancy_rate"),
            rooms=Count("id"),
        )
        .order_by("date")
    )
    rows = [
        {
            "date": r["date"].isoformat(),
            "base": money(r["base"]),
            "dynamic": money(r["dynamic"]),
            "competitor": money(r["competitor"]),
            "uplift": pct(float(r["dynamic"] or 0) - float(r["base"] or 0), r["base"]),
            "occupancy": money(r["occupancy"]),
            "rooms": r["rooms"],
        }
        for r in qs
    ]
    return (
        _cols(
            ("date", "Date"),
            ("base", "Avg base (€)"),
            ("dynamic", "Avg dynamic (€)"),
            ("competitor", "Avg competitor (€)"),
            ("uplift", "Uplift %"),
            ("occupancy", "Occupancy %"),
            ("rooms", "Rooms priced"),
        ),
        rows,
        {
            "Days": len(rows),
            "Avg uplift %": (
                round(sum(r["uplift"] for r in rows) / len(rows), 1) if rows else 0.0
            ),
        },
    )


def _executive(properties, report):
    ledger = NightLedger(properties, report.from_date, report.to_date)
    rows = ledger.per_property()
    k = ledger.kpis()
    return (
        _cols(
            ("name", "Property"),
            ("revenue", "Room revenue (€)"),
            ("occupancy", "Occupancy %"),
            ("adr", "ADR (€)"),
            ("revpar", "RevPAR (€)"),
        ),
        rows,
        {
            "Room revenue (€)": k["total_revenue"],
            "Occupancy %": k["occupancy_rate"],
            "ADR (€)": k["adr"],
            "RevPAR (€)": k["revpar"],
            "Bookings": k["bookings"],
        },
    )


BUILDERS = {
    "revenue": _revenue,
    "occupancy": _occupancy,
    "booking": _bookings,
    "payment": _payments,
    "guest": _guests,
    "pricing": _pricing,
    "custom": _executive,
}


def report_properties(report):
    if report.property_id:
        return [report.property]
    return list(Property.objects.order_by("name"))


def build_report(report):
    columns, rows, summary = BUILDERS[report.report_type](
        report_properties(report), report
    )
    return {"columns": columns, "rows": rows, "summary": summary}


def _csv(data):
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow([c["header"] for c in data["columns"]])
    for row in data["rows"]:
        writer.writerow([row.get(c["field"], "") for c in data["columns"]])
    return buffer.getvalue().encode("utf-8-sig"), "text/csv", "csv"


def _xlsx(report, data):
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = report.get_report_type_display()[:31]
    ws.append([c["header"] for c in data["columns"]])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in data["rows"]:
        ws.append([row.get(c["field"], "") for c in data["columns"]])
    summary = wb.create_sheet("Summary")
    for label, value in data["summary"].items():
        summary.append([label, value])
    buffer = io.BytesIO()
    wb.save(buffer)
    return (
        buffer.getvalue(),
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "xlsx",
    )


def _pdf(report, data):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import (
        Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle,
    )

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(A4), title=report.name)
    styles = getSampleStyleSheet()
    scope = report.property.name if report.property_id else "All properties"
    story = [
        Paragraph(report.name, styles["Title"]),
        Paragraph(
            f"{report.get_report_type_display()} · {scope} · {report.from_date} – {report.to_date}",
            styles["Normal"],
        ),
        Spacer(1, 12),
        Paragraph(
            " · ".join(f"<b>{k}</b>: {v}" for k, v in data["summary"].items()),
            styles["Normal"],
        ),
        Spacer(1, 12),
    ]
    table_rows = [[c["header"] for c in data["columns"]]] + [
        [str(row.get(c["field"], "")) for c in data["columns"]] for row in data["rows"]
    ]
    table = Table(table_rows, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1565C0")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTSIZE", (0, 0), (-1, -1), 7),
                (
                    "ROWBACKGROUNDS",
                    (0, 1),
                    (-1, -1),
                    [colors.white, colors.HexColor("#F2F5FA")],
                ),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#C9D1DC")),
            ]
        )
    )
    story.append(table)
    doc.build(story)
    return buffer.getvalue(), "application/pdf", "pdf"


def render_report(report, data=None):
    """Return (content, content_type, extension); falls back to CSV if a library is missing."""
    data = data or build_report(report)
    try:
        if report.export_format == "excel":
            return _xlsx(report, data)
        if report.export_format == "pdf":
            return _pdf(report, data)
    except ImportError:
        pass
    return _csv(data)
